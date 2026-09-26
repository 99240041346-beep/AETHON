from __future__ import annotations

import hashlib
import threading
from dataclasses import dataclass
from uuid import uuid4

MAX_ATTACHMENT_BYTES = 2 * 1024 * 1024
MAX_TEXT_CHARS = 120_000

_ALLOWED = {
    "text/plain": "text",
    "text/markdown": "text",
    "text/csv": "text",
    "application/json": "text",
    "application/xml": "text",
    "text/xml": "text",
    "image/png": "image",
    "image/jpeg": "image",
    "image/webp": "image",
    "image/gif": "image",
    "application/pdf": "binary",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "binary",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "binary",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "binary",
}

@dataclass(frozen=True)
class Attachment:
    attachment_id: str
    owner_id: str
    filename: str
    media_type: str
    size: int
    sha256: str
    kind: str
    text: str | None = None

_store: dict[str, Attachment] = {}
_lock = threading.RLock()

def _clean_text(raw: bytes) -> str:
    text = raw.decode("utf-8-sig", errors="replace").replace("\x00", "")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n[attachment text truncated]"
    return text

def _extract_document(raw: bytes, media_type: str) -> str | None:
    try:
        if media_type == "application/pdf":
            from pypdf import PdfReader
            import io
            pages = PdfReader(io.BytesIO(raw)).pages
            return "\n\n".join((page.extract_text() or "") for page in pages)[:MAX_TEXT_CHARS]
        if media_type.endswith("wordprocessingml.document"):
            from docx import Document
            import io
            doc = Document(io.BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs)[:MAX_TEXT_CHARS]
        if media_type.endswith("spreadsheetml.sheet"):
            from openpyxl import load_workbook
            import io
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join("" if value is None else str(value) for value in row))
                    if sum(len(x) for x in rows) >= MAX_TEXT_CHARS:
                        break
            return "\n".join(rows)[:MAX_TEXT_CHARS]
    except Exception:
        return None
    return None

def create_attachment(*, owner_id: str, filename: str, media_type: str, raw: bytes) -> Attachment:
    if media_type not in _ALLOWED:
        raise ValueError("unsupported attachment type")
    if len(raw) > MAX_ATTACHMENT_BYTES:
        raise ValueError("attachment exceeds the 2 MiB limit")
    safe_name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not safe_name or len(safe_name) > 180:
        raise ValueError("invalid attachment filename")
    kind = _ALLOWED[media_type]
    if kind == "binary" and media_type in {"application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}:
        extracted = _extract_document(raw, media_type)
        kind = "document"
    else:
        extracted = None
    text = _clean_text(raw) if kind == "text" else _extract_document(raw, media_type)
    item = Attachment(str(uuid4()), owner_id, safe_name, media_type, len(raw),
                      hashlib.sha256(raw).hexdigest(), kind, text or extracted)
    with _lock:
        _store[item.attachment_id] = item
        if len(_store) > 500:
            for key in list(_store)[:100]:
                _store.pop(key, None)
    return item

def get_attachment(attachment_id: str, owner_id: str) -> Attachment | None:
    with _lock:
        item = _store.get(attachment_id)
    if item is None or item.owner_id != owner_id:
        return None
    return item

def attachment_context(ids: list[str], owner_id: str) -> tuple[str, list[str]]:
    parts: list[str] = []
    names: list[str] = []
    for attachment_id in ids[:5]:
        item = get_attachment(attachment_id, owner_id)
        if item is None:
            continue
        names.append(item.filename)
        if item.text is not None:
            parts.append(f"Attachment: {item.filename}\n{item.text}")
        else:
            parts.append(f"Attachment: {item.filename} ({item.media_type}, {item.size} bytes). Content is attached metadata only; no visual/document extraction is available yet.")
    return "\n\n".join(parts), names


def retrieve_attachment_context(ids: list[str], owner_id: str, query: str, *,
                               max_chunks: int = 6, chunk_chars: int = 1800) -> str:
    """Bounded owner-scoped lexical retrieval over extracted attachment text."""
    if max_chunks < 1 or chunk_chars < 200:
        raise ValueError("invalid retrieval bounds")
    terms = {token.casefold() for token in query.split() if len(token) > 2}
    candidates: list[tuple[int, str, str]] = []
    for attachment_id in ids[:5]:
        item = get_attachment(attachment_id, owner_id)
        if item is None or not item.text:
            continue
        chunks = [item.text[i:i + chunk_chars] for i in range(0, len(item.text), chunk_chars)]
        for index, chunk in enumerate(chunks):
            lowered = chunk.casefold()
            score = sum(lowered.count(term) for term in terms)
            if score:
                candidates.append((score, item.filename, f"[{item.filename}#{index + 1}] {chunk}"))
    candidates.sort(key=lambda row: (-row[0], row[1]))
    return "\n\n".join(row[2] for row in candidates[:max_chunks])
